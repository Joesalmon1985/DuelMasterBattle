from __future__ import annotations

import hashlib
from dataclasses import dataclass, asdict
from pathlib import Path
from typing import Iterable, List

import openpyxl

from .normalize import normalize_text, stable_hash, stable_id


EXPECTED_HEADERS = [
    "Cast ID",
    "Soap Inspiration",
    "Situation Shape",
    "Character",
    "Village Role",
    "Story Role",
    "Personality Summary",
    "Dialogue State",
    "Trigger / Condition",
    "Opening Line",
    "Gives Choice?",
    "Choice #",
    "Choice Prompt",
    "Option A",
    "Response to A",
    "Option B",
    "Response to B",
    "State Change A",
    "State Change B",
    "Normal End Line",
    "Tragic End Line",
]

DEAD_TRAGIC = "(Dead in this tragic branch — no post-tragedy dialogue.)"


@dataclass(frozen=True)
class SourceRow:
    row_id: str
    cast_id: str
    sheet: str
    row_number: int
    soap_inspiration: str
    situation_shape: str
    character: str
    village_role: str
    story_role: str
    personality_summary: str
    dialogue_state: str
    trigger_condition: str
    opening_line: str
    gives_choice: bool
    choice_number: int | None
    choice_prompt: str
    option_a: str
    response_a: str
    option_b: str
    response_b: str
    state_change_a: str
    state_change_b: str
    normal_end_line: str
    tragic_end_line: str

    @property
    def story_context(self) -> str:
        parts = [
            f"Situation: {self.situation_shape}",
            f"Character: {self.character}",
            f"Village role: {self.village_role}",
            f"Story role: {self.story_role}",
            f"Personality: {self.personality_summary}",
            f"State: {self.dialogue_state}",
            f"Trigger: {self.trigger_condition}",
        ]
        if self.gives_choice:
            parts.extend([
                f"Choice A: {self.option_a}",
                f"Consequence A: {self.state_change_a}",
                f"Choice B: {self.option_b}",
                f"Consequence B: {self.state_change_b}",
            ])
        return "\n".join(p for p in parts if not p.endswith(": "))


@dataclass(frozen=True)
class BeatInstance:
    instance_id: str
    beat_id: str
    cast_id: str
    source_row_id: str
    source_sheet: str
    source_row: int
    source_field: str
    beat_type: str
    source_text: str
    normalized_text: str
    text_hash: str
    character: str
    village_role: str
    story_role: str
    personality_summary: str
    story_context: str
    consequential: bool
    branch_a: str | None
    branch_b: str | None


def workbook_sha256(path: Path) -> str:
    h = hashlib.sha256()
    with Path(path).open("rb") as f:
        for chunk in iter(lambda: f.read(1024 * 1024), b""):
            h.update(chunk)
    return h.hexdigest()


def _value(row: tuple, idx: int) -> str:
    value = row[idx] if idx < len(row) else None
    return "" if value is None else str(value).strip()


def read_cast_rows(path: Path, cast_id: str, sheet_name: str = "Dialogue Matrix") -> List[SourceRow]:
    wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    try:
        if sheet_name not in wb.sheetnames:
            raise ValueError(f"Workbook is missing required sheet '{sheet_name}'")
        ws = wb[sheet_name]
        header = [str(c.value).strip() if c.value is not None else "" for c in next(ws.iter_rows(min_row=1, max_row=1))]
        if header[:len(EXPECTED_HEADERS)] != EXPECTED_HEADERS:
            differences = [
                f"column {i+1}: expected {exp!r}, got {header[i] if i < len(header) else '<missing>'!r}"
                for i, exp in enumerate(EXPECTED_HEADERS)
                if i >= len(header) or header[i] != exp
            ]
            raise ValueError("Dialogue Matrix schema mismatch: " + "; ".join(differences[:8]))

        result: List[SourceRow] = []
        for row_no, row in enumerate(ws.iter_rows(min_row=2, values_only=True), start=2):
            if _value(row, 0) != cast_id:
                continue
            choice_no = None
            raw_choice_no = row[11] if len(row) > 11 else None
            if raw_choice_no not in (None, ""):
                try:
                    choice_no = int(raw_choice_no)
                except (TypeError, ValueError):
                    raise ValueError(f"Invalid Choice # at {sheet_name}!{row_no}: {raw_choice_no!r}")
            gives_choice = _value(row, 10).lower() in {"yes", "y", "true", "1"}
            source = SourceRow(
                row_id=stable_id("ROW", cast_id, sheet_name, row_no),
                cast_id=cast_id,
                sheet=sheet_name,
                row_number=row_no,
                soap_inspiration=_value(row, 1),
                situation_shape=_value(row, 2),
                character=_value(row, 3),
                village_role=_value(row, 4),
                story_role=_value(row, 5),
                personality_summary=_value(row, 6),
                dialogue_state=_value(row, 7),
                trigger_condition=_value(row, 8),
                opening_line=_value(row, 9),
                gives_choice=gives_choice,
                choice_number=choice_no,
                choice_prompt=_value(row, 12),
                option_a=_value(row, 13),
                response_a=_value(row, 14),
                option_b=_value(row, 15),
                response_b=_value(row, 16),
                state_change_a=_value(row, 17),
                state_change_b=_value(row, 18),
                normal_end_line=_value(row, 19),
                tragic_end_line=_value(row, 20),
            )
            if not source.character:
                raise ValueError(f"Missing Character at {sheet_name}!{row_no}")
            result.append(source)
        if not result:
            raise ValueError(f"No Dialogue Matrix rows found for Cast ID {cast_id}")
        return result
    finally:
        wb.close()


def extract_beats(rows: Iterable[SourceRow]) -> List[BeatInstance]:
    instances: List[BeatInstance] = []
    order = 0

    def add(row: SourceRow, field: str, beat_type: str, text: str, consequential: bool = False):
        nonlocal order
        text = str(text or "").strip()
        if not text or text == DEAD_TRAGIC:
            return
        order += 1
        normalized = normalize_text(text)
        digest = stable_hash(text)
        beat_id = stable_id(
            "BEAT", digest, beat_type, row.character, row.village_role, row.story_role,
            row.personality_summary, row.story_context,
            row.option_a if consequential else "", row.option_b if consequential else "",
        )
        instance_id = stable_id("INST", row.cast_id, row.sheet, row.row_number, field)
        instances.append(BeatInstance(
            instance_id=instance_id,
            beat_id=beat_id,
            cast_id=row.cast_id,
            source_row_id=row.row_id,
            source_sheet=row.sheet,
            source_row=row.row_number,
            source_field=field,
            beat_type=beat_type,
            source_text=text,
            normalized_text=normalized,
            text_hash=digest,
            character=row.character,
            village_role=row.village_role,
            story_role=row.story_role,
            personality_summary=row.personality_summary,
            story_context=row.story_context,
            consequential=consequential,
            branch_a=row.option_a if consequential else None,
            branch_b=row.option_b if consequential else None,
        ))

    for row in rows:
        if row.gives_choice:
            add(row, "Choice Prompt", "DECISION", row.choice_prompt, consequential=True)
            add(row, "Response to A", "AFTER_A", row.response_a)
            add(row, "Response to B", "AFTER_B", row.response_b)
        else:
            add(row, "Opening Line", "OPENING", row.opening_line)
        add(row, "Normal End Line", "NORMAL_END", row.normal_end_line)
        add(row, "Tragic End Line", "TRAGIC_END", row.tragic_end_line)
    return instances


def row_as_dict(row: SourceRow) -> dict:
    return asdict(row)
